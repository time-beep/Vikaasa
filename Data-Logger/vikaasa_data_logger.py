import serial
import serial.tools.list_ports
from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path
import threading

# ============================================================
# CONFIGURATION
# ============================================================
BAUD = 115200
EXCEL_FILE = Path(__file__).with_suffix(".xlsx")

# Leave PORT = None to choose from detected serial ports.
PORT = None

# Latest complete sensor observation received from the microcontroller.
latest = {
    "temp": None,
    "hum": None,
    "ldr": None,
    "soil": None,
    "red": None,
    "nir": None,
    "ndvi": None,
}

running = True


def choose_port():
    ports = list(serial.tools.list_ports.comports())
    if not ports:
        raise RuntimeError(
            "No serial ports found. Connect the microcontroller by USB."
        )

    print("\nDetected ports:")
    for i, p in enumerate(ports):
        print(f"  [{i}] {p.device} - {p.description}")

    if len(ports) == 1:
        return ports[0].device

    while True:
        try:
            n = int(input("Select port number: "))
            if 0 <= n < len(ports):
                return ports[n].device
        except ValueError:
            pass
        print("Enter a valid number.")


def calculate_ndvi(red, nir):
    """
    Calculate NDVI from red and NIR values.

    Important: this calculation is scientifically meaningful as NDVI when
    red and NIR represent suitable reflectance measurements (or calibrated
    proxies). Raw sensor counts may need calibration before interpretation.
    """
    denominator = nir + red
    if denominator == 0:
        return None
    return (nir - red) / denominator


def serial_reader(ser):
    """
    Expected serial formats:

    DATA,temp,humidity,ldr,soil,red,nir
        -> NDVI is calculated by this script.

    DATA,temp,humidity,ldr,soil,red,nir,ndvi
        -> NDVI is read directly from the microcontroller.

    Example:
    DATA,25.4,61.2,1834,2450,1200,1800
    """
    global latest, running

    while running:
        try:
            line = ser.readline().decode(errors="ignore").strip()
            if not line:
                continue

            parts = line.split(",")
            if parts[0] != "DATA":
                continue

            if len(parts) == 7:
                temp = float(parts[1])
                hum = float(parts[2])
                ldr = int(parts[3])
                soil = int(parts[4])
                red = float(parts[5])
                nir = float(parts[6])
                ndvi = calculate_ndvi(red, nir)

            elif len(parts) == 8:
                temp = float(parts[1])
                hum = float(parts[2])
                ldr = int(parts[3])
                soil = int(parts[4])
                red = float(parts[5])
                nir = float(parts[6])
                ndvi = float(parts[7])

            else:
                continue

            # Replace the whole snapshot so the main thread sees values
            # from one complete serial DATA record.
            latest = {
                "temp": temp,
                "hum": hum,
                "ldr": ldr,
                "soil": soil,
                "red": red,
                "nir": nir,
                "ndvi": ndvi,
            }

        except (ValueError, UnicodeDecodeError):
            # Ignore malformed or incomplete serial lines.
            continue
        except serial.SerialException:
            running = False
            break


def main():
    global running

    port = PORT if PORT else choose_port()

    if not EXCEL_FILE.exists():
        raise RuntimeError(
            f"Excel file not found: {EXCEL_FILE}\n"
            "Place the workbook beside this Python script."
        )

    wb = load_workbook(EXCEL_FILE)
    if "Vikaasa Data" not in wb.sheetnames:
        wb.close()
        raise RuntimeError(
            "Worksheet 'Vikaasa Data' not found in the Excel workbook."
        )

    ws = wb["Vikaasa Data"]

    # Required order: the three new columns must be before the final
    # Condition / Characteristic column.
    expected_headers = [
        "Reading #",
        "Timestamp",
        "Temperature (°C)",
        "Humidity (%)",
        "LDR (raw)",
        "Soil Moisture (raw)",
        "RED LED (raw)",
        "NIR LED (raw)",
        "NDVI",
        "Condition / Characteristic",
    ]

    actual_headers = [ws.cell(row=1, column=i).value
                      for i in range(1, len(expected_headers) + 1)]

    if actual_headers != expected_headers:
        wb.close()
        raise RuntimeError(
            "The workbook headers do not match the required format.\n"
            "Expected columns in this exact order:\n"
            + " | ".join(expected_headers)
            + "\nUpdate the workbook using the provided updated template."
        )

    print(f"\nOpening {port} at {BAUD} baud...")
    ser = serial.Serial(port, BAUD, timeout=0.2)

    thread = threading.Thread(target=serial_reader, args=(ser,), daemon=True)
    thread.start()

    print("\n==============================================")
    print(" Vikaasa Sensor Data Logger")
    print("==============================================")
    print("Receiving sensor data continuously.")
    print("Required DATA format:")
    print("  DATA,temp,humidity,ldr,soil,red,nir")
    print("NDVI will be calculated from RED and NIR.")
    print("Alternatively, append an NDVI value as the final DATA field.")
    print("")
    print("To record an observation:")
    print("  1. Prepare the plant/condition to be observed.")
    print("  2. Type its ground-truth condition or characteristic.")
    print("  3. Press ENTER to save the latest complete sensor reading.")
    print("")
    print("Type q and press ENTER to quit.")
    print("==============================================\n")

    reading_number = max(0, ws.max_row - 1)

    try:
        while True:
            condition = input("Condition / Characteristic: ")

            if condition.strip().lower() == "q":
                break

            # Snapshot the latest complete record at the time Enter is pressed.
            reading = latest.copy()
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            required_values = (
                reading["temp"], reading["hum"], reading["ldr"],
                reading["soil"], reading["red"], reading["nir"],
                reading["ndvi"]
            )
            if any(value is None for value in required_values):
                print(
                    "No complete reading containing temperature, humidity, "
                    "LDR, soil, RED and NIR has been received yet. Try again.\n"
                )
                continue

            reading_number += 1
            ws.append([
                reading_number,
                timestamp,
                reading["temp"],
                reading["hum"],
                reading["ldr"],
                reading["soil"],
                reading["red"],
                reading["nir"],
                reading["ndvi"],
                condition.strip(),
            ])

            # Save immediately after every recorded observation.
            wb.save(EXCEL_FILE)

            ndvi_display = (
                f"{reading['ndvi']:.4f}"
                if reading["ndvi"] is not None else "N/A"
            )
            print(
                f"Saved #{reading_number}: "
                f"T={reading['temp']:.1f}°C, "
                f"RH={reading['hum']:.1f}%, "
                f"LDR={reading['ldr']}, "
                f"Soil={reading['soil']}, "
                f"RED={reading['red']}, "
                f"NIR={reading['nir']}, "
                f"NDVI={ndvi_display}, "
                f"Condition='{condition.strip()}'\n"
            )

    except KeyboardInterrupt:
        print("\nStopping...")
    finally:
        running = False
        ser.close()
        wb.save(EXCEL_FILE)
        wb.close()
        print(f"\nExcel file saved: {EXCEL_FILE}")


if __name__ == "__main__":
    main()
